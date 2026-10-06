#!/usr/bin/env python3
"""Export Realtor.ca listings to CSV, Excel (.xlsx) or JSONL.

Calls the yugenox/realtor-ca-property-scraper Actor on Apify for one or more Canadian cities,
then writes one row per active listing: MLS number, price, beds, baths, size, address,
coordinates, listing date, main photo, listing URL and the listing brokerage. The Actor returns
the newest listings first.

--append keeps one growing file and adds only MLS numbers it doesn't have yet, so a daily run
tells you which listings are new since the last one.

Run with --help for every option. Requires Python 3.10+ and `pip install -r requirements.txt`
(openpyxl is only needed for .xlsx). APIFY_TOKEN comes from the environment or a .env file next
to this script; --from needs no token.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
import re
import sys
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

ACTOR = "yugenox/realtor-ca-property-scraper"

# The cities the Actor searches (its input schema, checked 2026-10-06). Anything else is
# skipped by the Actor, so the script stops before starting a run for it. --any-city skips
# this check, for when the Actor adds cities.
CITIES = [
    "Toronto", "Mississauga", "Brampton", "Vaughan", "Markham", "Oakville", "Hamilton",
    "Ottawa", "London", "Vancouver", "Surrey", "Burnaby", "Victoria", "Calgary", "Edmonton",
    "Montreal", "Winnipeg", "Halifax",
]
MAX_PER_CITY = 500  # the most listings one Realtor.ca search area returns
PRICE_PER_ROW = 0.001  # USD per listing (Apify Store, checked 2026-10-06), for the cost-cap warning

# Column order for CSV, XLSX and JSONL.
COLUMNS = [
    "mls_number", "price", "price_value", "bedrooms", "bathrooms", "size_interior",
    "property_type", "ownership_type", "stories", "parking", "address", "city_province", "postal_code",
    "province", "latitude", "longitude", "listed_date", "days_on_realtor", "photo_url", "listing_url",
    "brokerages", "listing_agents", "agent_phones", "searched", "first_seen",
]
NUMBER_COLUMNS = {"price_value", "latitude", "longitude", "days_on_realtor"}

EPILOG = f"""\
supported cities: {", ".join(CITIES)}
(no postal codes, provinces or neighbourhoods; write Montreal without the accent)

price: $0.001 per listing plus a start fee per run (Apify Store, checked 2026-10-06; the Store
page shows the current price). Not affiliated with CREA or REALTOR.ca.

examples:
  python listings.py "Toronto, ON" --limit 200 --out toronto.xlsx
  python listings.py "Vancouver, BC" --type rent --out vancouver-rentals.csv
  python listings.py --file cities.example.txt --append --out listings.csv --new-out new-today.csv
"""


def load_dotenv(path: Path = Path(__file__).with_name(".env")) -> None:
    """Minimal .env reader so the script has no extra dependency."""
    if not path.exists():
        return
    for line in path.read_text().splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def _field(obj, camel: str, snake: str):
    """apify-client v1/v2 return dicts, v3 returns model objects. Support both."""
    if obj is None:
        return None
    return obj.get(camel) if isinstance(obj, dict) else getattr(obj, snake, None)


def read_list(path: str) -> list[str]:
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    return [s.strip() for s in lines if s.strip() and not s.strip().startswith("#")]


def today() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d")


def join(values) -> str | None:
    parts = list(dict.fromkeys(v for v in (values or []) if v is not None and v != ""))
    return " | ".join(str(p) for p in parts) if parts else None


def listed_date(v) -> str | None:
    """Realtor.ca sends the listing date as .NET ticks (100 ns since 0001-01-01 UTC) or as an ISO
    string. Either way the export gets an ISO date-time in UTC."""
    if v is None or v == "":
        return None
    s = str(v)
    if re.fullmatch(r"\d{15,}", s):
        ms = int(s) // 10000 - 62135596800000
        return datetime.fromtimestamp(ms // 1000, tz=timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    try:
        d = datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        return s
    if d.tzinfo is None:
        d = d.replace(tzinfo=timezone.utc)
    return d.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def days_since(iso: str | None) -> int | None:
    """Whole days between listed_date and now. Computed here because the Actor's daysOnRealtor
    holds the number from Realtor.ca's "8 min ago" / "3 hours ago" label for listings under a day old."""
    if not iso:
        return None
    try:
        d = datetime.fromisoformat(iso.replace("Z", "+00:00"))
    except ValueError:
        return None
    return max(0, math.floor((datetime.now(timezone.utc) - d).total_seconds() / 86400))


def is_listing(r: dict) -> bool:
    return "mls_number" in r or "mlsNumber" in r


def to_row(r: dict, no_agents: bool) -> dict:
    """One Actor row (or one row this script wrote earlier) -> one flat output row."""
    if "mls_number" in r:
        row = {c: r.get(c) for c in COLUMNS}
    else:
        agents = r.get("agents") or []
        when = listed_date(r.get("listedDateUTC"))
        row = {
            "mls_number": r.get("mlsNumber"),
            "price": r.get("price"),
            "price_value": r.get("priceValue"),
            "bedrooms": r.get("bedrooms"),
            "bathrooms": r.get("bathrooms"),
            "size_interior": r.get("sizeInterior"),
            "property_type": r.get("propertyType"),
            "ownership_type": r.get("ownershipType"),
            "stories": r.get("stories"),
            "parking": r.get("parking"),
            "address": r.get("address"),
            "city_province": r.get("cityProvince"),
            "postal_code": r.get("postalCode"),
            "province": r.get("province"),
            "latitude": r.get("latitude"),
            "longitude": r.get("longitude"),
            "listed_date": when,
            "days_on_realtor": days_since(when),
            "photo_url": r.get("photoUrl"),
            "listing_url": r.get("listingUrl"),
            "brokerages": join(a.get("brokerage") for a in agents),
            "listing_agents": join(a.get("name") for a in agents),
            "agent_phones": join(a.get("phone") for a in agents),
            "searched": r.get("location"),
            "first_seen": today(),
        }
    if no_agents:
        row["listing_agents"] = None
        row["agent_phones"] = None
    if not row.get("first_seen"):
        row["first_seen"] = today()
    return row


def csv_value(v):
    """Text that starts with = + - @, a tab or a carriage return gets a leading ' so Excel and Google
    Sheets show it as text instead of running it as a formula."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def existing_ids(path: str) -> set[str]:
    p = Path(path)
    if not p.exists() or p.stat().st_size == 0:
        return set()
    if path.lower().endswith(".jsonl"):
        with p.open(encoding="utf-8-sig") as f:
            return {json.loads(line).get("mls_number") for line in f if line.strip()} - {None}
    with p.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames != COLUMNS:
            sys.exit(f"{path} has different columns than this script writes; use a new --out file")
        return {row["mls_number"] for row in reader if row.get("mls_number")}


def write_xlsx(path: str, rows: list[dict]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("Writing .xlsx needs openpyxl: pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "Listings"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([None] * len(COLUMNS))
        line = ws.max_row
        for i, col in enumerate(COLUMNS, start=1):
            v = r.get(col)
            if v is None or v == "":
                continue
            fmt = None
            if col in NUMBER_COLUMNS:
                v = float(v) if isinstance(v, str) else v
            elif col == "listed_date":
                # Excel has no time zones: the value is UTC.
                v = datetime.fromisoformat(str(v).replace("Z", "+00:00")).replace(tzinfo=None)
                fmt = "yyyy-mm-dd hh:mm"
            elif col == "first_seen":
                v = date.fromisoformat(str(v))
                fmt = "yyyy-mm-dd"
            else:
                v = ILLEGAL_CHARACTERS_RE.sub("", str(v))
            cell = ws.cell(row=line, column=i, value=v)
            if isinstance(v, str):
                cell.data_type = "s"  # never a formula, even if the text starts with =
            if fmt:
                cell.number_format = fmt
    widths = {"address": 30, "city_province": 40, "listing_url": 50, "photo_url": 50, "brokerages": 40,
              "listing_agents": 30, "agent_phones": 26, "listed_date": 17}
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 14)
    ws.freeze_panes = "A2"
    wb.save(path)


def write_rows(path: str, rows: list[dict], append: bool) -> None:
    low = path.lower()
    if low.endswith(".xlsx"):
        write_xlsx(path, rows)
        return
    p = Path(path)
    fresh = not append or not p.exists() or p.stat().st_size == 0
    if low.endswith(".jsonl"):
        with p.open("w" if fresh else "a", encoding="utf-8") as f:
            for r in rows:
                f.write(json.dumps(r, ensure_ascii=False, separators=(",", ":")) + "\n")
        return
    # UTF-8 with a BOM, so Excel shows accents (Montréal, Québec) correctly when it opens the file.
    with p.open("w" if fresh else "a", newline="", encoding="utf-8-sig" if fresh else "utf-8") as f:
        w = csv.writer(f)
        if fresh:
            w.writerow(COLUMNS)
        for r in rows:
            w.writerow([csv_value(r.get(c)) for c in COLUMNS])


def read_from_file(path: str) -> list[dict]:
    text = Path(path).read_text(encoding="utf-8-sig").strip()
    if not text:
        return []
    if text.startswith("["):
        return json.loads(text)
    return [json.loads(line) for line in text.splitlines() if line.strip()]


def main() -> None:
    p = argparse.ArgumentParser(
        prog="listings.py",
        description="Export Realtor.ca listings to CSV, Excel (.xlsx) or JSONL.",
        epilog=EPILOG,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    p.add_argument("cities", nargs="*", metavar="CITY", help='"Toronto, ON" or "Toronto" (see the list below)')
    p.add_argument("--file", metavar="PATH", help="text file with one city per line (# lines skipped)")
    s = p.add_argument_group("what to search")
    s.add_argument("--type", choices=["sale", "rent"], help="for sale (default) or for rent")
    s.add_argument("--property", choices=["residential", "commercial"], help="residential (default) or commercial")
    s.add_argument("--limit", type=int, default=100, metavar="N", help=f"listings per city (default 100, max {MAX_PER_CITY})")
    s.add_argument("--any-city", action="store_true", help="don't check cities against the list below")
    o = p.add_argument_group("output")
    o.add_argument("--out", default="listings.csv", metavar="PATH", help="output path: .csv (default listings.csv), .xlsx or .jsonl")
    o.add_argument("--append", action="store_true", help="add only listings whose mls_number is not in --out yet (.csv or .jsonl)")
    o.add_argument("--new-out", metavar="PATH", help="with --append: also write just this run's new listings here, replacing the file")
    o.add_argument("--no-agents", action="store_true", help="leave the listing agents' names and phone numbers out (the brokerage stays)")
    o.add_argument("--max-charge", type=float, default=2.0, metavar="USD", help="hard cost cap for the run in USD, enforced by Apify (default 2)")
    r = p.add_argument_group("re-export without a new run")
    r.add_argument("--dataset", metavar="ID", help="export the dataset of a run you already made (no new run starts)")
    r.add_argument("--from", dest="from_file", metavar="FILE", help="convert a JSON/JSONL file of the Actor's rows or of this script's rows (no token needed)")
    args = p.parse_args()

    cities = list(args.cities) + (read_list(args.file) if args.file else [])
    if args.limit < 1:
        p.error("--limit needs a whole number of 1 or more")
    if args.limit > MAX_PER_CITY:
        print(f"--limit {args.limit} -> {MAX_PER_CITY}: one Realtor.ca search area returns at most {MAX_PER_CITY} listings.", file=sys.stderr)
        args.limit = MAX_PER_CITY
    for f in (args.out, args.new_out):
        if f and not re.search(r"\.(csv|xlsx|jsonl)$", f, re.I):
            p.error(f"{f}: output files must end in .csv, .xlsx or .jsonl")
    if args.append and args.out.lower().endswith(".xlsx"):
        p.error("--append works with .csv or .jsonl (Excel opens the CSV)")
    if args.new_out and not args.append:
        p.error("--new-out needs --append (new = not in --out yet)")
    sources = sum(bool(x) for x in (cities, args.dataset, args.from_file))
    if sources == 0:
        p.error("give one or more cities, --file cities.txt, --dataset <id> or --from <file>")
    if sources > 1:
        p.error("use one of: cities/--file, --dataset, --from")
    if cities and not args.any_city:
        known = {c.lower() for c in CITIES}
        unknown = [c for c in cities if c.split(",")[0].strip().lower() not in known]
        if unknown:
            sys.exit(f"Not a supported city: {'; '.join(unknown)}\nSupported: {', '.join(CITIES)} "
                     "(no postal codes, provinces or neighbourhoods; write Montreal without the accent).")

    failed_run = False
    if args.from_file:
        items = read_from_file(args.from_file)
    else:
        load_dotenv()
        token = os.environ.get("APIFY_TOKEN")
        if not token:
            sys.exit("Set APIFY_TOKEN (see .env.example), or use --from <file>")
        from apify_client import ApifyClient

        client = ApifyClient(token)
        dataset_id = args.dataset
        if not dataset_id:
            run_input = {
                "locations": cities,
                "limitPerLocation": args.limit,
                "transactionType": args.type or "sale",
                "propertyType": args.property or "residential",
            }
            most = len(cities) * args.limit
            if most * PRICE_PER_ROW > args.max_charge:
                print(f"Up to {most} listings at ${PRICE_PER_ROW:g} each can cost more than the ${args.max_charge:g} cost cap, "
                      "and listings past the cap are not collected. Raise --max-charge to get them all.", file=sys.stderr)
            print(f"Running {ACTOR}: {len(cities)} city(ies), up to {args.limit} {run_input['propertyType']} "
                  f"listings for {run_input['transactionType']} each, cost cap ${args.max_charge:g} ...", file=sys.stderr)
            # max_total_charge_usd caps what this run can cost you (pay-per-event Actors only).
            run = client.actor(ACTOR).call(
                run_input=run_input, max_total_charge_usd=Decimal(str(args.max_charge))
            )
            if run is None:
                sys.exit("The run did not finish")
            status = _field(run, "status", "status")
            status = getattr(status, "value", status)
            message = _field(run, "statusMessage", "status_message")
            print(f"Run {_field(run, 'id', 'id')} finished: {status}" + (f" ({message})" if message else ""), file=sys.stderr)
            failed_run = status != "SUCCEEDED"
            dataset_id = _field(run, "defaultDatasetId", "default_dataset_id")
        items = list(client.dataset(dataset_id).iterate_items())
    if failed_run and not items:
        sys.exit(f"No listings, so {args.out} was left as it was.")

    rows = [to_row(it, args.no_agents) for it in items if is_listing(it)]
    skipped = 0
    if args.append:
        seen = existing_ids(args.out)
        everything = rows
        rows = [r for r in everything if not r.get("mls_number") or str(r["mls_number"]) not in seen]
        skipped = len(everything) - len(rows)
        # The Actor returns the newest listings first. If a city's run came back full and none of its
        # listings were in the file yet, it may not reach back to the previous run: some listings in
        # between were never seen.
        if seen and not args.from_file and not args.dataset:
            for city in dict.fromkeys(r.get("searched") for r in everything):
                got = [r for r in everything if r.get("searched") == city]
                if len(got) >= args.limit and all(str(r.get("mls_number")) not in seen for r in got):
                    print(f"{city}: all {len(got)} listings are new, so this run may not reach back to the "
                          "last one. Run more often or raise --limit to close the gap.", file=sys.stderr)
    write_rows(args.out, rows, args.append)
    if args.new_out:
        write_rows(args.new_out, rows, False)

    if args.append:
        extra = f" ({skipped} already there)" if skipped else ""
        also = f"; the new ones are also in {args.new_out}" if args.new_out else ""
        print(f"{len(rows)} new listings appended to {args.out}{extra}{also}", file=sys.stderr)
    else:
        print(f"Wrote {len(rows)} listings to {args.out}", file=sys.stderr)


if __name__ == "__main__":
    main()
