#!/usr/bin/env python3
"""Export the Realtor.ca agent directory to CSV, Excel (.xlsx) or JSONL.

Calls the yugenox/realtor-ca-agent-scraper Actor on Apify for one or more Canadian cities,
neighbourhoods or provinces, then writes one row per agent: name, title, brokerage, office
address, phone (when published), profile link and photo. Realtor.ca doesn't publish agent
emails, so there are none.

Run with --help for every option. Requires Python 3.10+ and `pip install -r requirements.txt`
(openpyxl is only needed for .xlsx). APIFY_TOKEN comes from the environment or a .env file next
to this script; --from needs no token.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

ACTOR = "yugenox/realtor-ca-agent-scraper"

# Realtor.ca's agent search never resolves a postal code (full or first 3 characters), so the
# script stops before starting a run for one.
POSTAL_RE = re.compile(r"^[A-Za-z]\d[A-Za-z][ -]?(\d[A-Za-z]\d)?$")
PRICE_PER_ROW = 0.001  # USD per agent (Apify Store, checked 2026-10-06), for the cost-cap warning

# Column order for CSV, XLSX and JSONL.
COLUMNS = [
    "individual_id", "name", "position", "brokerage", "office_address", "phone",
    "profile_url", "photo_url", "searched", "first_seen",
]

EPILOG = """\
locations: a city ("Toronto, ON", "Mississauga"), a district or neighbourhood ("North York",
"Yorkville") or a province ("Ontario", "ON"). Spelling must match Realtor.ca. Postal codes
don't work. Provinces work but are slower than naming cities.

price: $0.001 per agent plus a start fee per run (Apify Store, checked 2026-10-06; the Store page
shows the current price). Not affiliated with CREA or REALTOR.ca.

examples:
  python agents.py "Oakville, ON" --limit 100 --out oakville-agents.xlsx
  python agents.py "North York" "Yorkville" --out agents.csv
  python agents.py --file locations.example.txt --append --out agents.csv
"""


def load_dotenv(path: Path = Path(__file__).with_name(".env")) -> None:
    """Minimal .env reader so the script has no extra dependency."""
    if not path.exists():
        return
    for line in path.read_text().splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def _field(obj, camel: str, snake: str):
    """apify-client v1/v2 return dicts, v3 returns model objects. Support both."""
    if obj is None:
        return None
    return obj.get(camel) if isinstance(obj, dict) else getattr(obj, snake, None)


def read_list(path: str) -> list[str]:
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    return [s.strip() for s in lines if s.strip() and not s.strip().startswith("#")]


def today() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d")


def is_agent(r: dict) -> bool:
    return "individual_id" in r or "individualId" in r


def to_row(r: dict) -> dict:
    """One Actor row (or one row this script wrote earlier) -> one flat output row."""
    if "individual_id" in r:
        row = {c: r.get(c) for c in COLUMNS}
    else:
        row = {
            "individual_id": str(r["individualId"]) if r.get("individualId") is not None else None,
            "name": r.get("name"),
            "position": r.get("position"),
            "brokerage": r.get("brokerage"),
            "office_address": r.get("officeAddress"),
            "phone": r.get("phone"),
            "profile_url": r.get("profileUrl"),
            "photo_url": r.get("photoUrl"),
            "searched": r.get("location"),
            "first_seen": today(),
        }
    if not row.get("first_seen"):
        row["first_seen"] = today()
    return row


def csv_value(v):
    """Text that starts with = + - @, a tab or a carriage return gets a leading ' so Excel and Google
    Sheets show it as text instead of running it as a formula."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def existing_ids(path: str) -> set[str]:
    p = Path(path)
    if not p.exists() or p.stat().st_size == 0:
        return set()
    if path.lower().endswith(".jsonl"):
        with p.open(encoding="utf-8-sig") as f:
            return {json.loads(line).get("individual_id") for line in f if line.strip()} - {None}
    with p.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames != COLUMNS:
            sys.exit(f"{path} has different columns than this script writes; use a new --out file")
        return {row["individual_id"] for row in reader if row.get("individual_id")}


def write_xlsx(path: str, rows: list[dict]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("Writing .xlsx needs openpyxl: pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "Agents"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([None] * len(COLUMNS))
        line = ws.max_row
        for i, col in enumerate(COLUMNS, start=1):
            v = r.get(col)
            if v is None or v == "":
                continue
            if col == "first_seen":
                cell = ws.cell(row=line, column=i, value=date.fromisoformat(str(v)))
                cell.number_format = "yyyy-mm-dd"
                continue
            cell = ws.cell(row=line, column=i, value=ILLEGAL_CHARACTERS_RE.sub("", str(v)))
            cell.data_type = "s"  # never a formula, even if the text starts with =
    widths = {"name": 26, "position": 22, "brokerage": 40, "office_address": 50, "phone": 14,
              "profile_url": 44, "photo_url": 44, "searched": 16}
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 12)
    ws.freeze_panes = "A2"
    wb.save(path)


def write_rows(path: str, rows: list[dict], append: bool) -> None:
    low = path.lower()
    if low.endswith(".xlsx"):
        write_xlsx(path, rows)
        return
    p = Path(path)
    fresh = not append or not p.exists() or p.stat().st_size == 0
    if low.endswith(".jsonl"):
        with p.open("w" if fresh else "a", encoding="utf-8") as f:
            for r in rows:
                f.write(json.dumps(r, ensure_ascii=False, separators=(",", ":")) + "\n")
        return
    # UTF-8 with a BOM, so Excel shows accents correctly when it opens the file.
    with p.open("w" if fresh else "a", newline="", encoding="utf-8-sig" if fresh else "utf-8") as f:
        w = csv.writer(f)
        if fresh:
            w.writerow(COLUMNS)
        for r in rows:
            w.writerow([csv_value(r.get(c)) for c in COLUMNS])


def read_from_file(path: str) -> list[dict]:
    text = Path(path).read_text(encoding="utf-8-sig").strip()
    if not text:
        return []
    if text.startswith("["):
        return json.loads(text)
    return [json.loads(line) for line in text.splitlines() if line.strip()]


def main() -> None:
    p = argparse.ArgumentParser(
        prog="agents.py",
        description="Export the Realtor.ca agent directory to CSV, Excel (.xlsx) or JSONL.",
        epilog=EPILOG,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    p.add_argument("locations", nargs="*", metavar="LOCATION", help="city, district, neighbourhood or province (see below)")
    p.add_argument("--file", metavar="PATH", help="text file with one location per line (# lines skipped)")
    p.add_argument("--limit", type=int, default=50, metavar="N", help="agents per location (default 50)")
    o = p.add_argument_group("output")
    o.add_argument("--out", default="agents.csv", metavar="PATH", help="output path: .csv (default agents.csv), .xlsx or .jsonl")
    o.add_argument("--append", action="store_true", help="add only agents whose individual_id is not in --out yet (.csv or .jsonl)")
    o.add_argument("--max-charge", type=float, default=2.0, metavar="USD", help="hard cost cap for the run in USD, enforced by Apify (default 2)")
    r = p.add_argument_group("re-export without a new run")
    r.add_argument("--dataset", metavar="ID", help="export the dataset of a run you already made (no new run starts)")
    r.add_argument("--from", dest="from_file", metavar="FILE", help="convert a JSON/JSONL file of the Actor's rows or of this script's rows (no token needed)")
    args = p.parse_args()

    locations = list(args.locations) + (read_list(args.file) if args.file else [])
    if args.limit < 1:
        p.error("--limit needs a whole number of 1 or more")
    if not re.search(r"\.(csv|xlsx|jsonl)$", args.out, re.I):
        p.error("--out must end in .csv, .xlsx or .jsonl")
    if args.append and args.out.lower().endswith(".xlsx"):
        p.error("--append works with .csv or .jsonl (Excel opens the CSV)")
    sources = sum(bool(x) for x in (locations, args.dataset, args.from_file))
    if sources == 0:
        p.error("give one or more locations, --file locations.txt, --dataset <id> or --from <file>")
    if sources > 1:
        p.error("use one of: locations/--file, --dataset, --from")
    postal = [loc for loc in locations if POSTAL_RE.match(loc.split(",")[0].strip())]
    if postal:
        sys.exit(f"Realtor.ca's agent search doesn't take postal codes: {'; '.join(postal)}. "
                 "Use the city, district or neighbourhood instead.")

    failed_run = False
    if args.from_file:
        items = read_from_file(args.from_file)
    else:
        load_dotenv()
        token = os.environ.get("APIFY_TOKEN")
        if not token:
            sys.exit("Set APIFY_TOKEN (see .env.example), or use --from <file>")
        from apify_client import ApifyClient

        client = ApifyClient(token)
        dataset_id = args.dataset
        if not dataset_id:
            run_input = {"locations": locations, "limitPerLocation": args.limit}
            most = len(locations) * args.limit
            if most * PRICE_PER_ROW > args.max_charge:
                print(f"Up to {most} agents at ${PRICE_PER_ROW:g} each can cost more than the ${args.max_charge:g} cost cap, "
                      "and agents past the cap are not collected. Raise --max-charge to get them all.", file=sys.stderr)
            print(f"Running {ACTOR}: {len(locations)} location(s), up to {args.limit} agents each, "
                  f"cost cap ${args.max_charge:g} ...", file=sys.stderr)
            # max_total_charge_usd caps what this run can cost you (pay-per-event Actors only).
            run = client.actor(ACTOR).call(
                run_input=run_input, max_total_charge_usd=Decimal(str(args.max_charge))
            )
            if run is None:
                sys.exit("The run did not finish")
            status = _field(run, "status", "status")
            status = getattr(status, "value", status)
            message = _field(run, "statusMessage", "status_message")
            print(f"Run {_field(run, 'id', 'id')} finished: {status}" + (f" ({message})" if message else ""), file=sys.stderr)
            failed_run = status != "SUCCEEDED"
            dataset_id = _field(run, "defaultDatasetId", "default_dataset_id")
        items = list(client.dataset(dataset_id).iterate_items())
    if failed_run and not items:
        sys.exit(f"No agents, so {args.out} was left as it was.")

    rows = [to_row(it) for it in items if is_agent(it)]
    skipped = 0
    if args.append:
        seen = existing_ids(args.out)
        before = len(rows)
        rows = [r for r in rows if not r.get("individual_id") or str(r["individual_id"]) not in seen]
        skipped = before - len(rows)
    write_rows(args.out, rows, args.append)

    verb = "Appended" if args.append else "Wrote"
    extra = f" ({skipped} already there)" if skipped else ""
    print(f"{verb} {len(rows)} agents to {args.out}{extra}", file=sys.stderr)


if __name__ == "__main__":
    main()
